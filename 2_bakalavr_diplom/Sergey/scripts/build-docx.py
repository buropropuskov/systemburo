#!/usr/bin/env python3
"""
Сборка .docx файла ВКР Сергея из Markdown-файлов.

Формат: ГОСТ + требования МИРЭА (бакалавриат)
Шрифт: Times New Roman, 14 пт (основной), 18/16/14 пт (заголовки)
Межстрочный интервал: 1.5
Поля: верхнее 20мм, нижнее 20мм, левое 30мм, правое 15мм
"""

import os
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn, nsdecls
from docx.shared import Cm, Mm, Pt, RGBColor
from docx.enum.text import WD_BREAK

BASE_DIR = Path("/home/washka/project/diplom/Sergey")
OUTPUT_PATH = BASE_DIR / "output" / "vkr_sergey.docx"

INPUT_FILES = [
    BASE_DIR / "chapters" / "introduction.md",
    BASE_DIR / "chapters" / "chapter1.md",
    BASE_DIR / "chapters" / "chapter2.md",
    BASE_DIR / "chapters" / "chapter3.md",
    BASE_DIR / "chapters" / "conclusion.md",
    BASE_DIR / "sources" / "references.md",
    BASE_DIR / "chapters" / "appendices.md",
]

FONT_NAME = "Times New Roman"
FONT_SIZE = Pt(14)
HEADING1_SIZE = Pt(18)
HEADING2_SIZE = Pt(16)
HEADING3_SIZE = Pt(14)
CAPTION_SIZE = Pt(12)
TABLE_FONT_SIZE = Pt(12)
CODE_FONT = "Courier New"
CODE_FONT_SIZE = Pt(10)
LINE_SPACING = 1.5
FIRST_LINE_INDENT = Cm(1.25)

UNNUMBERED_HEADINGS = {
    "АННОТАЦИЯ", "ANNOTATION", "СОДЕРЖАНИЕ", "ВВЕДЕНИЕ",
    "ЗАКЛЮЧЕНИЕ", "СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ", "ПРИЛОЖЕНИЯ",
}

IMAGE_DIR = BASE_DIR / "images"
FIGURE_MAP = {
    "1.1": "xss1.png",
    "1.2": "xss2.png",
    "1.3": "xss3.png",
    "1.4": "react.png",
    "1.5": "vue.png",
    "2.1": "uc1.png",
    "2.2": "uc2.png",
    "2.3": "uc3.png",
    "2.4": "sitemap.png",
    "3.4": "sequence.png",
    "3.5": "components.png",
    "А.1": "bp1.png",
    "А.2": "bp2.png",
    "А.3": "bp3.png",
    "А.4": "activity_swimlane.png",
}


# ---------------------------------------------------------------------------
# Helpers: XML-level formatting
# ---------------------------------------------------------------------------

def set_paragraph_spacing(paragraph, before=0, after=0, line_spacing=LINE_SPACING):
    pf = paragraph.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line_spacing


def set_run_font(run, name=FONT_NAME, size=FONT_SIZE, bold=False, italic=False,
                 color=None):
    run.font.name = name
    run.font.size = size
    run.font.bold = bold
    run.font.italic = italic
    if color:
        run.font.color.rgb = color
    r = run._element
    rPr = r.find(qn("w:rPr"))
    if rPr is None:
        rPr = OxmlElement("w:rPr")
        r.insert(0, rPr)
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.insert(0, rFonts)
    rFonts.set(qn("w:eastAsia"), name)
    rFonts.set(qn("w:cs"), name)


def add_page_break_before(paragraph):
    pPr = paragraph._element.get_or_add_pPr()
    page_break = OxmlElement("w:pageBreakBefore")
    page_break.set(qn("w:val"), "true")
    pPr.append(page_break)


def set_keep_with_next(paragraph):
    pPr = paragraph._element.get_or_add_pPr()
    kwn = OxmlElement("w:keepNext")
    kwn.set(qn("w:val"), "true")
    pPr.append(kwn)


def set_keep_lines(paragraph):
    pPr = paragraph._element.get_or_add_pPr()
    kl = OxmlElement("w:keepLines")
    kl.set(qn("w:val"), "true")
    pPr.append(kl)


def set_widow_control(paragraph):
    pPr = paragraph._element.get_or_add_pPr()
    wc = OxmlElement("w:widowControl")
    wc.set(qn("w:val"), "true")
    pPr.append(wc)


def set_paragraph_borders(paragraph, color="000000", size="4", space="4"):
    pPr = paragraph._element.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    for border_name in ("top", "left", "bottom", "right"):
        border = OxmlElement(f"w:{border_name}")
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), size)
        border.set(qn("w:space"), space)
        border.set(qn("w:color"), color)
        pBdr.append(border)
    pPr.append(pBdr)


def set_paragraph_shading(paragraph, color="FFFFFF"):
    pPr = paragraph._element.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    pPr.append(shd)


def set_cell_shading(cell, color="FFFFFF"):
    tc = cell._element
    tcPr = tc.find(qn("w:tcPr"))
    if tcPr is None:
        tcPr = OxmlElement("w:tcPr")
        tc.insert(0, tcPr)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    tcPr.append(shd)


def set_table_borders(table):
    tbl = table._tbl
    tblPr = tbl.find(qn("w:tblPr"))
    if tblPr is None:
        tblPr = OxmlElement("w:tblPr")
        tbl.insert(0, tblPr)
    tblBorders = OxmlElement("w:tblBorders")
    for border_name in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = OxmlElement(f"w:{border_name}")
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "4")
        border.set(qn("w:space"), "0")
        border.set(qn("w:color"), "000000")
        tblBorders.append(border)
    tblPr.append(tblBorders)


def add_page_number_footer(section):
    footer = section.footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.clear()

    run = p.add_run()
    fld_char_begin = OxmlElement("w:fldChar")
    fld_char_begin.set(qn("w:fldCharType"), "begin")
    run._element.append(fld_char_begin)

    run2 = p.add_run()
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    run2._element.append(instr_text)

    run3 = p.add_run()
    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")
    run3._element.append(fld_char_end)

    for r in [run, run2, run3]:
        set_run_font(r, size=Pt(14))


def add_toc(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    run = p.add_run("СОДЕРЖАНИЕ")
    set_run_font(run, size=HEADING1_SIZE, bold=True)
    set_paragraph_spacing(p, after=12, line_spacing=LINE_SPACING)

    p_toc = doc.add_paragraph()
    run_toc = p_toc.add_run()
    fld_char_begin = OxmlElement("w:fldChar")
    fld_char_begin.set(qn("w:fldCharType"), "begin")
    run_toc._element.append(fld_char_begin)

    run_instr = p_toc.add_run()
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = ' TOC \\o "1-3" \\h \\z \\u '
    run_instr._element.append(instr)

    run_sep = p_toc.add_run()
    fld_char_sep = OxmlElement("w:fldChar")
    fld_char_sep.set(qn("w:fldCharType"), "separate")
    run_sep._element.append(fld_char_sep)

    run_placeholder = p_toc.add_run(
        "Обновите оглавление: ПКМ -> Обновить поле (Ctrl+A, F9)"
    )
    set_run_font(run_placeholder, italic=True, color=RGBColor(128, 128, 128))

    run_end = p_toc.add_run()
    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")
    run_end._element.append(fld_char_end)

    set_paragraph_spacing(p_toc, line_spacing=LINE_SPACING)

    p_break = doc.add_paragraph()
    run_break = p_break.add_run()
    run_break.add_break(WD_BREAK.PAGE)


# ---------------------------------------------------------------------------
# Style setup
# ---------------------------------------------------------------------------

def setup_styles(doc):
    style = doc.styles["Normal"]
    font = style.font
    font.name = FONT_NAME
    font.size = FONT_SIZE
    pf = style.paragraph_format
    pf.line_spacing = LINE_SPACING
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.first_line_indent = FIRST_LINE_INDENT
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    _fix_style_font(style, FONT_NAME)

    h1 = doc.styles["Heading 1"]
    h1.font.name = FONT_NAME
    h1.font.size = HEADING1_SIZE
    h1.font.bold = True
    h1.font.color.rgb = RGBColor(0, 0, 0)
    h1.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    h1.paragraph_format.space_before = Pt(0)
    h1.paragraph_format.space_after = Pt(12)
    h1.paragraph_format.line_spacing = LINE_SPACING
    h1.paragraph_format.first_line_indent = FIRST_LINE_INDENT
    h1.paragraph_format.page_break_before = True
    h1.paragraph_format.keep_with_next = True
    _fix_heading_font(h1, FONT_NAME)

    h2 = doc.styles["Heading 2"]
    h2.font.name = FONT_NAME
    h2.font.size = HEADING2_SIZE
    h2.font.bold = True
    h2.font.color.rgb = RGBColor(0, 0, 0)
    h2.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    h2.paragraph_format.space_before = Pt(24)
    h2.paragraph_format.space_after = Pt(12)
    h2.paragraph_format.line_spacing = LINE_SPACING
    h2.paragraph_format.first_line_indent = FIRST_LINE_INDENT
    h2.paragraph_format.page_break_before = False
    h2.paragraph_format.keep_with_next = True
    h2.paragraph_format.keep_together = True
    h2.paragraph_format.widow_control = True
    _fix_heading_font(h2, FONT_NAME)

    h3 = doc.styles["Heading 3"]
    h3.font.name = FONT_NAME
    h3.font.size = HEADING3_SIZE
    h3.font.bold = True
    h3.font.color.rgb = RGBColor(0, 0, 0)
    h3.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    h3.paragraph_format.space_before = Pt(24)
    h3.paragraph_format.space_after = Pt(24)
    h3.paragraph_format.line_spacing = LINE_SPACING
    h3.paragraph_format.first_line_indent = FIRST_LINE_INDENT
    h3.paragraph_format.page_break_before = False
    h3.paragraph_format.keep_with_next = True
    _fix_heading_font(h3, FONT_NAME)


def _fix_heading_font(style, name):
    rPr = style.element.find(qn("w:rPr"))
    if rPr is None:
        rPr = OxmlElement("w:rPr")
        style.element.append(rPr)
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.insert(0, rFonts)
    rFonts.set(qn("w:ascii"), name)
    rFonts.set(qn("w:hAnsi"), name)
    rFonts.set(qn("w:eastAsia"), name)
    rFonts.set(qn("w:cs"), name)


def _fix_style_font(style, name):
    rPr = style.element.find(qn("w:rPr"))
    if rPr is None:
        rPr = OxmlElement("w:rPr")
        style.element.append(rPr)
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.insert(0, rFonts)
    rFonts.set(qn("w:eastAsia"), name)
    rFonts.set(qn("w:cs"), name)


def setup_page(doc):
    section = doc.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Mm(20)
    section.bottom_margin = Mm(20)
    section.left_margin = Mm(30)
    section.right_margin = Mm(15)
    add_page_number_footer(section)


# ---------------------------------------------------------------------------
# Markdown parser
# ---------------------------------------------------------------------------

RE_BOLD_ITALIC = re.compile(r"\*\*\*(.*?)\*\*\*")
RE_BOLD = re.compile(r"\*\*(.*?)\*\*")
RE_ITALIC = re.compile(r"(?<!\*)\*(?!\*)(.*?)(?<!\*)\*(?!\*)")
RE_INLINE_CODE = re.compile(r"`([^`]+)`")
RE_FIGURE_PLACEHOLDER = re.compile(
    r"^\[Рисунок\s+([\dА-Яа-я]+\.\d+)\s*[-—–]\s*(.+)\]$"
)
RE_FIGURE_BARE = re.compile(
    r"^Рисунок\s+([\dА-Яа-я]+\.\d+)\s*[-—–]\s*(.+)$"
)
RE_LISTING_CAPTION = re.compile(
    r"^(?:\*?Листинг\s+[\dА-Яа-я]+\.\d+\s*[-—–]\s*.+\*?)$"
)
RE_TABLE_ROW = re.compile(r"^\|(.+)\|$")
RE_TABLE_SEPARATOR = re.compile(r"^\|[-:|]+\|$")
RE_HEADING1_HASH = re.compile(r"^#\s+(.+)$")
RE_HEADING2_HASH = re.compile(r"^##\s+(.+)$")
RE_HEADING3_HASH = re.compile(r"^###\s+(.+)$")
RE_HEADING4_HASH = re.compile(r"^####\s+(.+)$")
RE_NUMBERED_LIST = re.compile(r"^(\d+)\.\s+(.+)$")
RE_BULLET_LIST = re.compile(r"^[-*]\s+(.+)$")
RE_BLOCKQUOTE = re.compile(r"^>\s*(.*)$")
RE_HORIZONTAL_RULE = re.compile(r"^---+$")
RE_MD_TABLE_CAPTION = re.compile(
    r"^Таблица\s+([\dА-Яа-я]+\.\d+)\s*[-—–]\s*(.+)$"
)


def parse_inline(paragraph, text, base_bold=False, base_italic=False,
                 font_size=None):
    segments = _split_inline(text)
    for seg_text, seg_bold, seg_italic, seg_code in segments:
        if not seg_text:
            continue
        run = paragraph.add_run(seg_text)
        if seg_code:
            set_run_font(run, name=CODE_FONT,
                         size=font_size or FONT_SIZE,
                         bold=base_bold, italic=base_italic)
        else:
            set_run_font(run, size=font_size or FONT_SIZE,
                         bold=seg_bold or base_bold,
                         italic=seg_italic or base_italic)


def _split_inline(text):
    result = []
    pos = 0
    length = len(text)

    while pos < length:
        next_match = None
        next_pos = length
        next_type = None

        m = RE_INLINE_CODE.search(text, pos)
        if m and m.start() < next_pos:
            next_match = m
            next_pos = m.start()
            next_type = "code"

        m = RE_BOLD_ITALIC.search(text, pos)
        if m and m.start() < next_pos:
            next_match = m
            next_pos = m.start()
            next_type = "bold_italic"

        m = RE_BOLD.search(text, pos)
        if m and m.start() < next_pos:
            next_match = m
            next_pos = m.start()
            next_type = "bold"

        m = RE_ITALIC.search(text, pos)
        if m and m.start() < next_pos:
            next_match = m
            next_pos = m.start()
            next_type = "italic"

        if next_match is None:
            result.append((text[pos:], False, False, False))
            break

        if next_pos > pos:
            result.append((text[pos:next_pos], False, False, False))

        inner = next_match.group(1)
        if next_type == "code":
            result.append((inner, False, False, True))
        elif next_type == "bold_italic":
            result.append((inner, True, True, False))
        elif next_type == "bold":
            result.append((inner, True, False, False))
        elif next_type == "italic":
            result.append((inner, False, True, False))

        pos = next_match.end()

    return result


def _detect_level_by_numbering(text):
    m = re.match(r"^(\d+(?:\.\d+)*)\s+", text)
    if m:
        num = m.group(1)
        dots = num.count(".")
        if dots == 0:
            return 1
        elif dots == 1:
            return 2
        else:
            return 3
    return None


def _is_unnumbered_heading(text):
    return text.strip().upper() in UNNUMBERED_HEADINGS


def is_section_heading(text):
    m = RE_HEADING1_HASH.match(text)
    if m:
        inner = m.group(1).strip()
        level = _detect_level_by_numbering(inner)
        centered = _is_unnumbered_heading(inner)
        return (level or 1, inner, centered)

    m = RE_HEADING2_HASH.match(text)
    if m:
        inner = m.group(1).strip()
        level = _detect_level_by_numbering(inner)
        centered = _is_unnumbered_heading(inner)
        return (level or 2, inner, centered)

    m = RE_HEADING3_HASH.match(text)
    if m:
        inner = m.group(1).strip()
        level = _detect_level_by_numbering(inner)
        centered = _is_unnumbered_heading(inner)
        return (level or 2, inner, centered)

    m = RE_HEADING4_HASH.match(text)
    if m:
        inner = m.group(1).strip()
        level = _detect_level_by_numbering(inner)
        return (level or 3, inner, False)

    return None


class MarkdownToDocx:

    def __init__(self, doc):
        self.doc = doc
        self.in_code_block = False
        self.code_lines = []
        self.code_lang = ""
        self.in_table = False
        self.table_rows = []
        self.is_first_heading = True

    def process_file(self, filepath, is_references=False):
        with open(filepath, "r", encoding="utf-8") as f:
            lines = f.readlines()

        if is_references:
            lines = self._filter_references(lines)

        i = 0
        while i < len(lines):
            line = lines[i].rstrip("\n")
            i += 1

            if line.startswith("```"):
                if not self.in_code_block:
                    self._flush_table()
                    self.in_code_block = True
                    self.code_lang = line[3:].strip()
                    self.code_lines = []
                else:
                    self._flush_code_block()
                continue

            if self.in_code_block:
                self.code_lines.append(line)
                continue

            if RE_BLOCKQUOTE.match(line):
                continue

            if RE_HORIZONTAL_RULE.match(line):
                continue

            if not line.strip():
                self._flush_table()
                continue

            if RE_TABLE_ROW.match(line):
                if RE_TABLE_SEPARATOR.match(line):
                    continue
                self.table_rows.append(line)
                continue
            else:
                self._flush_table()

            heading = is_section_heading(line)
            if heading:
                level, text, centered = heading
                self._add_heading(level, text, centered)
                continue

            # Listing caption (italic)
            stripped = line.strip()
            if stripped.startswith("*Листинг") and stripped.endswith("*"):
                caption_text = stripped.strip("*").strip()
                self._add_listing_caption(caption_text)
                continue

            m = RE_LISTING_CAPTION.match(stripped)
            if m:
                self._add_listing_caption(stripped)
                continue

            m = RE_MD_TABLE_CAPTION.match(stripped)
            if m:
                self._add_table_caption(m.group(0))
                continue

            m = RE_FIGURE_PLACEHOLDER.match(stripped)
            if m:
                self._add_figure_caption(m.group(0).strip("[]"))
                continue

            m = RE_FIGURE_BARE.match(stripped)
            if m:
                self._add_figure_caption(m.group(0))
                continue

            m = RE_NUMBERED_LIST.match(stripped)
            if m:
                self._add_numbered_item(m.group(1), m.group(2))
                continue

            m = RE_BULLET_LIST.match(stripped)
            if m:
                self._add_bullet_item(m.group(1))
                continue

            self._add_paragraph(stripped)

        self._flush_table()
        if self.in_code_block:
            self._flush_code_block()

    def _filter_references(self, lines):
        result = []
        for line in lines:
            stripped = line.strip()
            if stripped.startswith("#"):
                result.append(line)
            elif stripped:
                result.append(line)
        return result

    def _add_heading(self, level, text, centered=False):
        if level == 1:
            display_text = text.upper()
            p = self.doc.add_heading(display_text, level=1)

            if self.is_first_heading:
                p.paragraph_format.page_break_before = False
                self.is_first_heading = False

            if centered:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.first_line_indent = Cm(0)
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                p.paragraph_format.first_line_indent = FIRST_LINE_INDENT

            for run in p.runs:
                set_run_font(run, size=HEADING1_SIZE, bold=True)
            set_keep_with_next(p)

        elif level == 2:
            p = self.doc.add_heading(text, level=2)
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            for run in p.runs:
                set_run_font(run, size=HEADING2_SIZE, bold=True)
            set_keep_with_next(p)
            set_keep_lines(p)
            set_widow_control(p)

        elif level == 3:
            p = self.doc.add_heading(text, level=3)
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            for run in p.runs:
                set_run_font(run, size=HEADING3_SIZE, bold=True)
            set_keep_with_next(p)

    def _add_paragraph(self, text):
        p = self.doc.add_paragraph()
        p.style = self.doc.styles["Normal"]
        parse_inline(p, text)
        set_paragraph_spacing(p)

    def _add_numbered_item(self, number, text):
        p = self.doc.add_paragraph()
        p.style = self.doc.styles["Normal"]
        parse_inline(p, f"{number}. {text}")
        set_paragraph_spacing(p)

    def _add_bullet_item(self, text):
        p = self.doc.add_paragraph()
        p.style = self.doc.styles["Normal"]
        parse_inline(p, f"\u2013 {text}")
        set_paragraph_spacing(p)

    def _add_figure_caption(self, text):
        try:
            from PIL import Image as PILImage
            has_pil = True
        except ImportError:
            has_pil = False

        MAX_WIDTH_CM = 15.5
        MAX_HEIGHT_CM = 23.0

        m = re.match(r"Рисунок\s+([\dА-Яа-я]+\.\d+)", text)
        if m and has_pil:
            fig_num = m.group(1)
            if fig_num in FIGURE_MAP:
                img_path = IMAGE_DIR / FIGURE_MAP[fig_num]
                if img_path.exists():
                    with PILImage.open(img_path) as img:
                        w_px, h_px = img.size
                    aspect = h_px / w_px

                    width_cm = MAX_WIDTH_CM
                    height_cm = width_cm * aspect

                    if height_cm > MAX_HEIGHT_CM:
                        height_cm = MAX_HEIGHT_CM
                        width_cm = height_cm / aspect

                    img_paragraph = self.doc.add_paragraph()
                    img_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    img_paragraph.paragraph_format.first_line_indent = Cm(0)
                    set_paragraph_spacing(img_paragraph, before=6, after=0, line_spacing=1.0)
                    run = img_paragraph.add_run()
                    run.add_picture(str(img_path), width=Cm(width_cm))

        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Cm(0)
        run = p.add_run(text)
        set_run_font(run, size=CAPTION_SIZE, bold=True, italic=False)
        set_paragraph_spacing(p, before=0, after=6, line_spacing=1.0)
        set_widow_control(p)

    def _add_listing_caption(self, text):
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.first_line_indent = Cm(0)
        run = p.add_run(text)
        set_run_font(run, size=CAPTION_SIZE, bold=False, italic=True)
        set_paragraph_spacing(p, before=6, after=0, line_spacing=1.0)
        set_keep_with_next(p)

    def _add_table_caption(self, text):
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.first_line_indent = Cm(0)
        run = p.add_run(text)
        set_run_font(run, size=CAPTION_SIZE, bold=False, italic=True)
        set_paragraph_spacing(p, before=6, after=0, line_spacing=1.0)
        set_keep_with_next(p)

    def _flush_code_block(self):
        self.in_code_block = False
        if not self.code_lines:
            return

        while self.code_lines and not self.code_lines[0].strip():
            self.code_lines.pop(0)
        while self.code_lines and not self.code_lines[-1].strip():
            self.code_lines.pop()

        for idx, code_line in enumerate(self.code_lines):
            p = self.doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.left_indent = Cm(0.5)
            set_paragraph_spacing(p, line_spacing=1.0)
            set_paragraph_borders(p)
            set_paragraph_shading(p, "FFFFFF")
            run = p.add_run(code_line if code_line else " ")
            set_run_font(run, name=CODE_FONT, size=CODE_FONT_SIZE)

        p_after = self.doc.add_paragraph()
        p_after.style = self.doc.styles["Normal"]
        set_paragraph_spacing(p_after, before=6, line_spacing=LINE_SPACING)

        self.code_lines = []

    def _flush_table(self):
        if not self.table_rows:
            return

        parsed_rows = []
        for row_line in self.table_rows:
            cells = [c.strip() for c in row_line.strip("|").split("|")]
            parsed_rows.append(cells)

        if not parsed_rows:
            self.table_rows = []
            return

        num_cols = len(parsed_rows[0])
        num_rows = len(parsed_rows)

        table = self.doc.add_table(rows=num_rows, cols=num_cols)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.style = self.doc.styles["Table Grid"]
        set_table_borders(table)

        tbl = table._tbl
        tblPr = tbl.tblPr if tbl.tblPr is not None else OxmlElement("w:tblPr")
        tbl.insert(0, tblPr)
        tblLayout = OxmlElement("w:tblLayout")
        tblLayout.set(qn("w:type"), "autofit")
        tblPr.append(tblLayout)

        for r_idx, row_data in enumerate(parsed_rows):
            for c_idx, cell_text in enumerate(row_data):
                if c_idx >= num_cols:
                    break
                cell = table.cell(r_idx, c_idx)
                p = cell.paragraphs[0]
                p.clear()
                if r_idx == 0:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    parse_inline(p, cell_text, base_bold=True,
                                 font_size=TABLE_FONT_SIZE)
                else:
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    parse_inline(p, cell_text, font_size=TABLE_FONT_SIZE)

                p.paragraph_format.first_line_indent = Cm(0)
                p.paragraph_format.space_before = Pt(1)
                p.paragraph_format.space_after = Pt(1)
                set_paragraph_spacing(p, line_spacing=1.0)
                set_cell_shading(cell, "FFFFFF")

        p_after = self.doc.add_paragraph()
        p_after.style = self.doc.styles["Normal"]
        set_paragraph_spacing(p_after, before=6, line_spacing=LINE_SPACING)

        self.table_rows = []


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def build_docx():
    doc = Document()

    setup_styles(doc)
    setup_page(doc)

    converter = MarkdownToDocx(doc)

    # 1. Содержание (TOC)
    add_toc(doc)

    # 2. Все файлы
    for filepath in INPUT_FILES:
        is_ref = filepath.name == "references.md"
        print(f"  Обработка: {filepath.name}")
        converter.process_file(filepath, is_references=is_ref)

    # Сохранение
    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUTPUT_PATH))

    file_size = OUTPUT_PATH.stat().st_size
    file_size_kb = file_size / 1024
    total_chars = 0
    for fp in INPUT_FILES:
        with open(fp, "r", encoding="utf-8") as f:
            total_chars += len(f.read())
    estimated_pages = max(1, total_chars // 2800)

    print()
    print(f"  Файл:    {OUTPUT_PATH}")
    print(f"  Размер:  {file_size_kb:.1f} КБ ({file_size} байт)")
    print(f"  Страниц: ~{estimated_pages} (оценка по объёму текста)")
    print()
    print("  Для обновления оглавления откройте файл в Word")
    print("  и нажмите Ctrl+A, затем F9.")


if __name__ == "__main__":
    print()
    print("Сборка ВКР Сергея в .docx...")
    print()

    missing = [f for f in INPUT_FILES if not f.exists()]
    if missing:
        print("ОШИБКА: не найдены файлы:")
        for f in missing:
            print(f"  {f}")
        sys.exit(1)

    build_docx()
    print("  Готово.")
