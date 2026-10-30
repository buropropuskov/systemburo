"""Проверки отказов, которые раньше могли дать ложный успешный выпуск."""
from pathlib import Path
import subprocess
import json
import hashlib
import tempfile
import unittest
from unittest.mock import patch
from contextlib import ExitStack, redirect_stdout
from io import StringIO

import doc_tools
import build_docs
import docs_impact
import doc_facts


class CommandChecks(unittest.TestCase):
    def test_shared_detail_component_requires_role_review(self):
        shared = 'frontend/src/components/CreateApplication/EmployeeDetailsModal.vue'
        hits = docs_impact.impact([shared])
        documents = {hit[0] for hit in hits}
        self.assertIn('Ролевые руководства', documents)
        self.assertNotIn('Руководство пользователя', documents)

    def test_application_form_still_maps_to_applicant_document(self):
        path = 'frontend/src/components/CreateApplication/CreateApplication.vue'
        documents = {hit[0] for hit in docs_impact.impact([path])}
        self.assertIn('Руководство пользователя', documents)

    def test_recent_security_supply_and_draft_files_are_mapped(self):
        paths = ['.env.example', 'go.mod', 'frontend/package-lock.json',
                 '.github/scripts/go_test_groups.py',
                 'internal/services/session_access.go',
                 'internal/middleware/file_access.go',
                 'internal/services/document_access_policy.go',
                 'internal/services/application_file_quota.go',
                 'frontend/src/components/OpenPassagesModal.vue']
        self.assertEqual(docs_impact.uncovered(paths), [])
        for path in paths:
            with self.subTest(path=path):
                self.assertTrue(docs_impact.impact([path]))

    def test_removed_environment_parameter_is_reported(self):
        reference = '# Приложение Б\n| `REMOVED_OPTION` | x |\n# Приложение В\n'
        with patch.object(doc_facts, 'env_defaults', return_value={}), \
                patch.object(doc_facts, 'operational_env_names', return_value=set()):
            problems = []
            doc_facts.check_env(reference, problems)
        self.assertEqual(len(problems), 1)
        self.assertIn('REMOVED_OPTION', problems[0])

    def test_operational_parameter_is_not_required_in_go_config(self):
        reference = '# Приложение Б\n| `EDGE_OPTION` | x |\n# Приложение В\n'
        with patch.object(doc_facts, 'env_defaults', return_value={}), \
                patch.object(doc_facts, 'operational_env_names', return_value={'EDGE_OPTION'}):
            problems = []
            doc_facts.check_env(reference, problems)
        self.assertEqual(problems, [])

    def test_operational_environment_names_cover_supply_sources(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            (root / 'scripts').mkdir()
            (root / 'nginx').mkdir()
            (root / 'frontend/src/api').mkdir(parents=True)
            (root / '.env.example').write_text('TEMPLATE_OPTION=\n', encoding='utf-8')
            (root / 'docker-compose.base.yml').write_text(
                '    EDGE_OPTION: ${SOURCE_OPTION:-}\n', encoding='utf-8')
            (root / 'scripts/env.sh').write_text('echo "$SCRIPT_OPTION"\n', encoding='utf-8')
            (root / 'nginx/entrypoint.sh').write_text('echo "${PROXY_OPTION:-}"\n', encoding='utf-8')
            (root / 'frontend/src/api/client.js').write_text(
                'const value = import.meta.env.VITE_OPTION;\n', encoding='utf-8')
            with patch.object(doc_facts, 'REPO', directory):
                names = doc_facts.operational_env_names()
        self.assertEqual(names, {'TEMPLATE_OPTION', 'EDGE_OPTION', 'SOURCE_OPTION',
                                 'SCRIPT_OPTION', 'PROXY_OPTION', 'VITE_OPTION'})

    def test_environment_duplicates_only_within_reference_table(self):
        reference = '# Приложение Б\n| `MODE` | x |\n# Приложение В\n'
        with patch.object(doc_facts, 'env_defaults', return_value={'MODE': 'x'}):
            problems = []
            doc_facts.check_env('| `MODE` | x |\n' + reference, problems)
            self.assertEqual(problems, [])
            doc_facts.check_env(reference.replace('# Приложение В', '| `MODE` | x |\n# Приложение В'), problems)
            self.assertEqual(len(problems), 1)

    def test_git_failure_is_not_an_empty_result(self):
        with tempfile.TemporaryDirectory() as directory:
            with self.assertRaises(RuntimeError):
                doc_tools.git(directory, 'rev-parse', '--verify', 'refs/heads/nonexistent-docs-check')

    def test_missing_source_fact_is_not_zero(self):
        with self.assertRaises(ValueError):
            doc_tools.number(r'ttl:\s*(\d+)', 'константа удалена')
        with self.assertRaises(ValueError):
            doc_tools.section('func Other() {}', r'^func AllModels')

    def test_model_list_survives_source_line_shift(self):
        code = '\n' * 250 + 'func AllModels() []interface{} {\n' \
               '\treturn []interface{}{\n\t\t&models.One{},\n\t\t&models.Two{},\n\t}\n}\n'
        part = doc_tools.section(code, r'^func AllModels\(\) \[\]interface\{\} \{')
        self.assertIn('&models.Two{}', part)

    def test_sql_tables_and_models_are_not_counted_twice(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            (root / 'internal/database').mkdir(parents=True)
            (root / 'internal/models').mkdir(parents=True)
            (root / 'internal/database/partitions.go').write_text(
                'CREATE TABLE request_logs (\nCREATE TABLE pd_audit_logs (', encoding='utf-8')
            (root / 'internal/models/logs.go').write_text(
                'func (r *RequestLogs) TableName() string { return "request_logs" }', encoding='utf-8')
            self.assertEqual(doc_tools.schema_table_count(directory, '&models.RequestLogs{}'), 2)

    def test_staged_and_untracked_names_with_spaces_are_preserved(self):
        with tempfile.TemporaryDirectory() as directory:
            doc_tools.git(directory, 'init', '-q')
            path = Path(directory) / 'new file.go'
            path.write_text('package example\n', encoding='utf-8')
            with patch.object(docs_impact, 'REPO', directory):
                self.assertEqual(docs_impact.changed(''), ['new file.go'])
                doc_tools.git(directory, 'add', '--', path.name)
                self.assertEqual(docs_impact.changed(''), ['new file.go'])


class PdfChecks(unittest.TestCase):
    @staticmethod
    def transfer_table_fixture():
        # Normalized layout fragment from deploy p163. In the complete page,
        # '.env в' ends at 1109, 'Сертификат' at 1172, but the first-column
        # alternative 'Файл параметров' appears only in the later callout (1277).
        text = (
            'Таблица 14.1 — Состав переносимых данных '
            'Элемент Где находится Последствия, если не перенести '
            'База данных Том pgdata Утрата всех заявок, реестров, учётных записей и журналов '
            'Загруженные Том uploads Фотографии, шаблоны бланков и приложения к '
            'файлы заявкам пропадут, записи в базе будут ссылаться в пустоту '
            'Файл .env в каталоге См. предупреждение ниже параметров системы '
            'Сертификат nginx/certs/ Потребуется выпустить заново, что при смене '
            'адреса и так необходимо '
            'ОПАСНО. Файл параметров переносится со старого сервера, а не создаётся заново.'
        )
        headers = ['Элемент', 'Где находится', 'Последствия, если не перенести']
        rows = [
            ['База данных', 'Том pgdata', 'Утрата всех заявок'],
            ['Загруженные файлы', 'Том uploads', 'Фотографии, шаблоны бланков'],
            ['Файл параметров', '.env в каталоге системы', 'См. предупреждение ниже'],
            ['Сертификат', 'nginx/certs/', 'Потребуется выпустить заново'],
        ]
        metadata = [{'number': '14.1',
                     'head_probes': [build_docs.cell_probe(h) for h in headers],
                     'row_probes': build_docs.table_probes(headers, rows)}]
        return text, metadata

    def test_complete_table_does_not_borrow_probe_from_later_callout(self):
        text, metadata = self.transfer_table_fixture()
        with patch.object(build_docs, 'executable', return_value='pdftotext'), \
             patch.object(build_docs.subprocess, 'run', return_value=
                          subprocess.CompletedProcess([], 0, text, '')):
            self.assertEqual(build_docs.find_first_unsplit_break({}, metadata, set()),
                             (None, None))

    def test_real_table_break_is_still_found_before_missing_last_row(self):
        text, metadata = self.transfer_table_fixture()
        text = text.replace(
            'Сертификат nginx/certs/ Потребуется выпустить заново, что при смене '
            'адреса и так необходимо ', '') + '\fСертификат nginx/certs/ Потребуется'
        with patch.object(build_docs, 'executable', return_value='pdftotext'), \
             patch.object(build_docs.subprocess, 'run', return_value=
                          subprocess.CompletedProcess([], 0, text, '')):
            self.assertEqual(build_docs.find_first_unsplit_break({}, metadata, set()),
                             ('14.1', [3]))

    def test_tail_uses_earliest_alternative_before_real_next_break(self):
        text, metadata = self.transfer_table_fixture()
        text = text.replace('Таблица 14.1 — Состав переносимых данных',
                            'Продолжение таблицы 14.1')
        metadata[0]['row_probes'] = ([['previous']] + metadata[0]['row_probes']
                                     + [['next page']])
        with patch.object(build_docs, 'executable', return_value='pdftotext'), \
             patch.object(build_docs.subprocess, 'run', return_value=
                          subprocess.CompletedProcess([], 0, text + '\fnext page', '')):
            self.assertEqual(build_docs.next_break_in_tail({}, metadata, '14.1', [1]), 5)

    def test_row_probe_selection_is_independent_of_alternative_order(self):
        for probes, expected in [(['later', 'early'], 5), (['early', 'later'], 5),
                                 (['missing', 'early'], 5), (['missing'], -1),
                                 ([], -1)]:
            with self.subTest(probes=probes):
                self.assertEqual(build_docs.find_row_probe('early then later', probes, 0), expected)
        self.assertEqual(build_docs.find_row_probe('early then later', ['early', 'later'], 5), 16)
        self.assertEqual(build_docs.find_row_probe('15m 5m', ['5m'], 0), 6)

    def test_table_caption_is_not_an_unlabelled_continuation(self):
        for caption in ['Таблица 9.14 — Поля, которые затирает команда',
                        'Таблица Б.9 — Параметры уведомлений']:
            with self.subTest(caption=caption):
                self.assertFalse(build_docs.starts_with_table_header(
                    [caption, 'Таблица Что затирается', 'Сотрудники Имя'],
                    {'Таблица Что затирается'}))

    def test_real_table_column_named_table_still_requires_continuation(self):
        self.assertTrue(build_docs.starts_with_table_header(
            ['Таблица Что затирается', 'Сотрудники Имя'],
            {'Таблица Что затирается'}))

    def test_page_map_ignores_contents_entry_without_dotted_leader(self):
        text = ('СОДЕРЖАНИЕ\n1. Назначение ...... 3\n'
                '2. Длинный заголовок4\f'
                '3\n1. Назначение\nОсновной текст\f'
                '4\n2. Длинный заголовок\nПродолжение\f')
        with patch.object(build_docs, 'executable', return_value='pdftotext'), \
             patch.object(build_docs.subprocess, 'run', return_value=
                          subprocess.CompletedProcess([], 0, text, '')):
            self.assertEqual(build_docs.pagemap_from_pdf({}, [
                (1, '1. Назначение'), (1, '2. Длинный заголовок')]),
                {'1. Назначение': 2, '2. Длинный заголовок': 3})

    def test_page_map_rejects_document_without_body_heading(self):
        with patch.object(build_docs, 'executable', return_value='pdftotext'), \
             patch.object(build_docs.subprocess, 'run', return_value=
                          subprocess.CompletedProcess([], 0, 'СОДЕРЖАНИЕ\n1. Назначение ... 3', '')):
            with self.assertRaisesRegex(RuntimeError, 'первый заголовок'):
                build_docs.pagemap_from_pdf({}, [(1, '1. Назначение')])

    def test_tail_break_must_advance_past_existing_split(self):
        metadata = [{'number': '1', 'head_probes': [],
                     'row_probes': [['first'], ['second'], ['third']]}]
        for text, expected in [('Продолжение таблицы 1 missing', None),
                               ('Продолжение таблицы 1 second', 2),
                               ('Продолжение таблицы 1 second third', None)]:
            with self.subTest(text=text), \
                 patch.object(build_docs, 'executable', return_value='pdftotext'), \
                 patch.object(build_docs.subprocess, 'run', return_value=
                              subprocess.CompletedProcess([], 0, text, '')):
                self.assertEqual(build_docs.next_break_in_tail({}, metadata, '1', [1]), expected)

    def test_tail_search_ignores_tables_outside_reported_silent_pages(self):
        metadata = [{'number': '1', 'head_probes': [],
                     'row_probes': [['first'], ['second'], ['third']]}]
        text = 'Продолжение таблицы 1 second\fHeader third\fOther page'
        with patch.object(build_docs, 'executable', return_value='pdftotext'), \
             patch.object(build_docs.subprocess, 'run', return_value=
                          subprocess.CompletedProcess([], 0, text, '')):
            self.assertIsNone(build_docs.next_break_in_tail(
                {}, metadata, '1', [1], silent_pages=[3]))
            self.assertEqual(build_docs.next_break_in_tail(
                {}, metadata, '1', [1], silent_pages=[2]), 2)

    def test_empty_changelog_form_is_not_a_data_continuation(self):
        from docx import Document
        document = Document()
        with patch.object(build_docs, 'changelog', return_value=[
            ('1.0', '01.01.2026', 'First'), ('1.1', '02.01.2026', 'Second')
        ]):
            build_docs.change_log(document, {}, {'И.1': [1]})
        captions = [paragraph.text for paragraph in document.paragraphs
                    if 'таблицы И.1' in paragraph.text]
        self.assertEqual(captions, ['Продолжение таблицы И.1', 'Окончание таблицы И.1'])

    def test_changelog_page_breaks_belong_to_content_not_empty_paragraphs(self):
        from docx import Document
        document = Document()
        with patch.object(build_docs, 'changelog', return_value=[
            ('1.0', '01.01.2026', 'First')
        ]):
            build_docs.change_log(document, {})
        self.assertFalse(document.element.xpath('.//w:br[@w:type="page"]'))
        heading = next(p for p in document.paragraphs
                       if p.text == 'ЛИСТ РЕГИСТРАЦИИ ИЗМЕНЕНИЙ')
        ending = next(p for p in document.paragraphs
                      if p.text == 'Окончание таблицы И.1')
        self.assertTrue(heading.paragraph_format.page_break_before)
        self.assertTrue(ending.paragraph_format.page_break_before)

    def test_changelog_ending_has_no_empty_spacer_before_its_page_break(self):
        from docx import Document
        document = Document()
        with patch.object(build_docs, 'changelog', return_value=[
            ('1.0', '01.01.2026', 'First')
        ]):
            build_docs.change_log(document, {})
        ending = next(p for p in document.paragraphs
                      if p.text == 'Окончание таблицы И.1')
        self.assertTrue(ending._p.getprevious().tag.endswith('}tbl'))

    def test_wrapped_table_header_after_running_notice_is_detected(self):
        rows = build_docs.page_body_lines('Не для распространения\n№ Показатель Допустимое Фактически\nзначение\n27 Доступность\n8\n')
        self.assertTrue(build_docs.starts_with_table_header(
            rows, {'№ Показатель Допустимое значение Фактически'}))
        self.assertFalse(build_docs.starts_with_table_header(
            ['Продолжение таблицы 4.2', *rows], {'№ Показатель Допустимое значение Фактически'}))

    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.root = Path(self.directory.name)
        self.pdf = self.root / 'Проверка.pdf'
        self.pdf.write_bytes(b'%PDF-old')
        (self.root / 'Проверка.docx').write_bytes(b'fixture')
        (self.root / 'Проверка.md').write_text('# Проверка\n', encoding='utf-8')
        self.paths = patch.object(build_docs, 'doc_path', side_effect=lambda cfg, ext: str(self.root / ('Проверка' + ext)))
        self.paths.start()
        self.addCleanup(self.paths.stop)

    def test_failed_conversion_does_not_accept_existing_pdf(self):
        failure = subprocess.CompletedProcess([], 1, '', 'conversion failed')
        with patch.object(build_docs.subprocess, 'run', return_value=failure):
            with self.assertRaises(RuntimeError):
                build_docs.to_pdf({})
        self.assertEqual(self.pdf.read_bytes(), b'%PDF-old')

    def test_zero_exit_without_new_pdf_is_failure(self):
        with patch.object(build_docs.subprocess, 'run', return_value=subprocess.CompletedProcess([], 0, '', '')):
            with self.assertRaises(RuntimeError):
                build_docs.to_pdf({})
        self.assertEqual(self.pdf.read_bytes(), b'%PDF-old')

    def test_invalid_output_does_not_replace_existing_pdf(self):
        def convert(command, **kwargs):
            (Path(command[command.index('--outdir') + 1]) / self.pdf.name).write_bytes(b'not a PDF')
            return subprocess.CompletedProcess(command, 0, '', '')
        with patch.object(build_docs.subprocess, 'run', side_effect=convert):
            with self.assertRaises(RuntimeError):
                build_docs.to_pdf({})
        self.assertEqual(self.pdf.read_bytes(), b'%PDF-old')

    def test_success_replaces_pdf_with_fresh_output(self):
        def convert(command, **kwargs):
            (Path(command[command.index('--outdir') + 1]) / self.pdf.name).write_bytes(b'%PDF-new')
            return subprocess.CompletedProcess(command, 0, '', '')
        with patch.object(build_docs.subprocess, 'run', side_effect=convert):
            self.assertEqual(build_docs.to_pdf({}), str(self.pdf))
        self.assertEqual(self.pdf.read_bytes(), b'%PDF-new')

    def test_profile_is_outside_document_output_directory(self):
        def convert(command, **kwargs):
            output = Path(command[command.index('--outdir') + 1])
            profile = next(arg for arg in command if arg.startswith('-env:UserInstallation='))
            self.assertFalse(profile.split('=', 1)[1].startswith(output.as_uri() + '/'))
            (output / self.pdf.name).write_bytes(b'%PDF-new')
            return subprocess.CompletedProcess(command, 0, '', '')
        with patch.object(build_docs.subprocess, 'run', side_effect=convert):
            build_docs.to_pdf({})

    def test_manifest_records_exact_output_hashes(self):
        with patch.object(build_docs, 'git', return_value=subprocess.CompletedProcess([], 0, 'revision\n', '')):
            build_docs.write_build_manifest({'key': 'pilot'}, 14)
        record = json.loads((self.root / 'Проверка.build.json').read_text(encoding='utf-8'))
        self.assertEqual(record['code_revision'], 'revision')
        self.assertEqual(record['pdf_sha256'], hashlib.sha256(b'%PDF-old').hexdigest())
        self.assertEqual(record['pages'], 14)

    def test_final_layout_can_fix_many_tables_without_immediate_count_reduction(self):
        metadata = [{'number': str(n), 'row_probes': [['first'], ['second']]}
                    for n in range(8)]
        repaired = []
        searches = []
        def find_break(*args):
            searches.append(True)
            return (None, None) if len(searches) == 1 else (str(len(repaired)), [1])
        def verdict(cfg, number, **kwargs):
            repaired.append(number)
            return 'ok'
        def report(*args, **kwargs):
            # Сдвиг следующей таблицы оставляет число разрывов прежним.
            return ([1], []) if len(repaired) < 8 else ([], [])
        with ExitStack() as stack:
            stack.enter_context(patch.object(build_docs, 'build', return_value=([(1, 'Title')], metadata)))
            stack.enter_context(patch.object(build_docs, 'to_pdf', return_value=str(self.pdf)))
            stack.enter_context(patch.object(build_docs, 'find_first_unsplit_break', side_effect=find_break))
            stack.enter_context(patch.object(build_docs, 'split_verdict', side_effect=verdict))
            stack.enter_context(patch.object(build_docs, 'report_silent_breaks', side_effect=report))
            stack.enter_context(patch.object(build_docs, 'pagemap_from_pdf', return_value={'Title': 1}))
            stack.enter_context(patch.object(build_docs.subprocess, 'run', return_value=
                subprocess.CompletedProcess([], 0, b'Pages: 1', b'')))
            manifest = stack.enter_context(patch.object(build_docs, 'write_build_manifest'))
            stack.enter_context(redirect_stdout(StringIO()))
            build_docs.build_one('deploy')
        self.assertEqual(repaired, [str(n) for n in range(8)])
        manifest.assert_called_once()

    def test_unresolved_final_break_cannot_publish_a_manifest(self):
        with ExitStack() as stack:
            stack.enter_context(patch.object(build_docs, 'build', return_value=([(1, 'Title')], [])))
            stack.enter_context(patch.object(build_docs, 'to_pdf', return_value=str(self.pdf)))
            stack.enter_context(patch.object(build_docs, 'find_first_unsplit_break', return_value=(None, None)))
            stack.enter_context(patch.object(build_docs, 'report_silent_breaks', return_value=([3], [])))
            stack.enter_context(patch.object(build_docs, 'pagemap_from_pdf', return_value={'Title': 1}))
            manifest = stack.enter_context(patch.object(build_docs, 'write_build_manifest'))
            stack.enter_context(redirect_stdout(StringIO()))
            with self.assertRaisesRegex(RuntimeError, 'разрывы таблиц'):
                build_docs.build_one('deploy')
        manifest.assert_not_called()


if __name__ == '__main__':
    unittest.main()
